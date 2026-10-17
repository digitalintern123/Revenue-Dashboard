"""upload_checks (pre-save checks) and format_registry (what is in the file)."""

import datetime as dt
import io

import openpyxl
import pandas as pd

from modules import format_registry as fr
from modules import import_report
from modules import upload_checks as uc


def _keys(rows):
    return pd.DataFrame(rows, columns=["date", "location", "outlet", "pax", "revenue"])


def test_copied_day_is_flagged():
    outlets = [f"O{i}" for i in range(6)]
    rows = []
    for d in ("2026-04-04", "2026-04-05"):
        rows += [(d, "Delhi", o, 10 + i, 1000 + i) for i, o in enumerate(outlets)]
    rows += [("2026-04-06", "Delhi", o, 20 + i, 2000 + i) for i, o in enumerate(outlets)]
    c = uc.copied_days(_keys(rows))
    assert c.level == uc.WARN
    assert c.table.iloc[0]["Dates with identical figures"] == "2026-04-04, 2026-04-05"


def test_pax_without_revenue():
    rows = [(f"2026-05-{d:02d}", "Delhi", "Reserved Lounge", 400, 0) for d in range(1, 8)]
    rows += [(f"2026-05-{d:02d}", "Delhi", "Porter", 400, 5000) for d in range(1, 8)]
    c = uc.pax_without_revenue(_keys(rows))
    assert c.level == uc.WARN and list(c.table["Outlet"]) == ["Reserved Lounge"]


def test_new_outlet_with_spelling_suggestion():
    stored = pd.DataFrame({"location": ["Delhi"], "outlet": ["International Lounge"]})
    f = _keys([("2026-05-01", "Delhi", "Internationl Lounge", 1, 10),
               ("2026-05-01", "Delhi", "International Lounge", 1, 10)])
    c = uc.new_outlets(f, stored)
    assert c.count == 1 and c.table.iloc[0]["Possibly the same as"] == "International Lounge"


def test_changes_vs_stored_counts_new_same_and_different():
    stored = _keys([("2026-05-01", "Delhi", "A", 10, 100), ("2026-05-01", "Delhi", "B", 10, 100)])
    f = _keys([("2026-05-01", "Delhi", "A", 10, 100), ("2026-05-01", "Delhi", "B", 10, 150),
               ("2026-05-02", "Delhi", "A", 10, 100)])
    c = uc.changes_vs_stored(f, stored)
    assert c.level == uc.WARN and c.count == 1
    assert "1 new line(s), 1 already stored" in c.detail


def _long_sheet(extra_revenue_row=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Date", "Location", "Business", "Sub-Business", "PAX", "Revenue"])
    ws.append([dt.date(2026, 5, 1), "Delhi", "Lounges", "INL 5&6", 10, 1000])
    ws.append([None, "Delhi", "Lounges", "INL 5&6", 5, 500])            # not imported
    if extra_revenue_row:
        ws.append([None, None, None, None, None, 77])                     # parser never sees it
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_reconcile_accounts_for_rows_not_imported():
    skipped = pd.DataFrame([{"Result": import_report.NOT_IMPORTED, "PAX": "5", "Revenue": "500"}])
    ok = uc.reconcile_long_sheet(_long_sheet(), "Data", 0, 1000, 10, skipped)
    assert ok.level == uc.OK
    bad = uc.reconcile_long_sheet(_long_sheet(True), "Data", 0, 1000, 10, skipped)
    assert bad.level == uc.WARN and "Unexplained" in bad.detail


def _traffic(rows):
    return pd.DataFrame(rows, columns=["date", "location", "terminal", "granularity", "traffic"])


def test_traffic_copied_and_unusual_days():
    rows = []
    for d in range(1, 11):
        v = 5000 + d * 10
        rows += [(f"2026-04-{d:02d}", "Delhi", "T1 Dep", "daily", v),
                 (f"2026-04-{d:02d}", "Delhi", "T2 Dep", "daily", v + 1)]
    rows[-2:] = [("2026-04-10", "Delhi", "T1 Dep", "daily", 5050),        # = 5 Apr
                 ("2026-04-10", "Delhi", "T2 Dep", "daily", 5051)]
    rows.append(("2026-04-11", "Delhi", "T1 Dep", "daily", 1000))          # far below usual
    t = _traffic(rows)
    assert uc.traffic_copied_days(t).table.iloc[0]["Dates with identical traffic"] == "2026-04-05, 2026-04-10"
    unusual = uc.traffic_unusual_days(t)
    assert unusual.count == 1 and unusual.table.iloc[0]["Date"] == "2026-04-11"


def test_traffic_month_vs_stored_previous_month():
    f = _traffic([("2026-05-01", "Hyderabad", "Domestic", "daily", 30_000)])
    prev = pd.DataFrame({"date": [dt.date(2026, 4, 1)], "location": ["Hyderabad"],
                         "terminal": ["Domestic"], "traffic": [60_000]})
    c = uc.traffic_vs_previous_month(f, prev)
    assert c.level == uc.WARN and c.table.iloc[0]["Change"] == "-50%"


def test_aop_months_already_stored():
    rows = pd.DataFrame({"location": ["Delhi"] * 2, "segment": ["EHPL"] * 2,
                         "business_unit": ["Lounges"] * 2, "outlet": ["INL"] * 2,
                         "year": [2026, 2026], "month": [4, 5], "aop": [1e5, 2e5]})
    a = uc.analyse_aop({"aop_rows": rows, "skipped_rows": pd.DataFrame(),
                        "units_multiplier": 100000, "format": "outlet_monthly"}, None, [(2026, 4)])
    assert a.checks[0].title == "1 of 2 month(s) already stored"


def test_identify_revenue_lists_every_sheet():
    from modules import data_processor
    data = _long_sheet()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb.create_sheet("AOP").append(["x"])
    wb.create_sheet("Blank")
    buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()
    res = data_processor.process_uploaded_file(io.BytesIO(data), "rev.xlsx", save_to_db=False)
    assert res.success and res.layout == data_processor.LAYOUT_LONG and res.source_sheet == "Data"
    t = fr.identify_revenue(data, "rev.xlsx", res, 1.0).set_index("Sheet")
    assert t.loc["Data", "Result"] == fr.IMPORTED and t.loc["Data", "Confidence"] == "100%"
    assert t.loc["AOP", "Result"] == fr.NOT_USED and "AOP Targets" in t.loc["AOP", "What the app found"]
    assert t.loc["Blank", "Result"] == fr.EMPTY


def test_identify_traffic_names_the_reader_per_sheet():
    from modules import traffic_parser as tp
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hyd"
    ws.append(["DOMESTIC PAX (Schedule)", None, None, "International PAX (Schedule)", None, None])
    ws.append([None, "ARRIVAL", "DEPARTURE", "TOTAL", "ARRIVAL", "DEPARTURE", "TOTAL"])
    for d in range(1, 4):
        ws.append([dt.datetime(2026, 9, d), 10, 20, 30, 1, 2, 3])
    wb.create_sheet("Delhi").append(["just a title"])
    buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()
    df = tp.parse_traffic_auto(io.BytesIO(data), "t.xlsx")
    t = fr.identify_traffic(data, "t.xlsx", df, 0.95).set_index("Sheet")
    assert t.loc["Hyd", "Result"] == fr.IMPORTED and "side by side" in t.loc["Hyd", "What the app found"]
    assert t.loc["Delhi", "Result"] == fr.NOT_RECOGNISED
