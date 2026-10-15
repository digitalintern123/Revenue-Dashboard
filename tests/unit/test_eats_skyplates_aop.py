"""Tests for the Sky Plates / Encalm Eats AOP importer (label-driven)."""

import datetime as dt
import io

import openpyxl
import pytest

from modules import eats_skyplates_aop_parser as esp


def _workbook(extra_rows=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anything"
    ws.cell(1, 2, "Revenue Analysis at Unit Level")
    ws.cell(1, 4, 100000)
    hdr = ["Geographical Segment", "Business Segment", "Unit-ID", "Sales Unit-ID"]
    for i, h in enumerate(hdr):
        ws.cell(2, 2 + i, h)
    months = [dt.datetime(2026, 4, 1), dt.datetime(2026, 5, 2), dt.datetime(2026, 8, 3)]
    for i, m in enumerate(months):
        ws.cell(2, 6 + i, m)

    r = 3 + extra_rows          # extra_rows shifts everything: parser must not care
    def put(geo, seg, unit, vals):
        nonlocal r
        ws.cell(r, 2, geo); ws.cell(r, 3, seg); ws.cell(r, 4, unit)
        for i, v in enumerate(vals):
            ws.cell(r, 6 + i, v)
        r += 1

    put("Delhi", "Lounge", "T2", [5, 5, 5])                 # EHPL row, ignored here
    r += 1
    put("Hyd", "IFK", "In-Flight Kitchen Hyd", [0, 1.5, 2])
    r += 1
    put("Delhi", "F&B", "T3 Arrivals F&B", [3, 3, 3])       # head = sum of children
    put(None, None, "Chaayos", [1, 1, 1])
    put(None, "some note", "CDH", [2, 2, 2])
    r += 1
    put("Delhi", "F&B", "T3 Departures F&B", [4, 4, 4])
    put(None, None, "Chaayos 54-55", [2, 2, 2])
    put(None, None, "Starbucks", [1, 1, 1])
    put(None, None, "Starbucks", [1, 1, 1])                 # duplicate name
    r += 1
    put("All", "F&B", "Events & ODCs", [1, 1, 1])
    put(None, None, "Events & ODCs", [1, 1, 1])
    r += 1
    put("Goa", "Hotel", "Goa Business Hotel", [9, 9, 9])    # closes section, ignored

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_label_driven_and_row_shift_invariant():
    a = esp.parse_eats_and_skyplates_aop(_workbook())["aop_rows"]
    b = esp.parse_eats_and_skyplates_aop(_workbook(extra_rows=4))["aop_rows"]
    assert a.sort_values(list(a.columns)).reset_index(drop=True).equals(
        b.sort_values(list(b.columns)).reset_index(drop=True))


def test_units_and_sky_plates():
    res = esp.parse_eats_and_skyplates_aop(_workbook())
    sp = res["aop_rows"][res["aop_rows"]["segment"] == "Sky Plates"]
    assert set(sp["location"]) == {"Hyderabad"}
    assert set(sp["outlet"]) == {"Encalm Sky Plates"}
    assert len(sp) == 2                                   # April is 0 -> not stored
    assert sp["aop"].sum() == pytest.approx(3.5 * 100000)


def test_children_only_no_double_count_and_names():
    res = esp.parse_eats_and_skyplates_aop(_workbook())
    ee = res["aop_rows"][res["aop_rows"]["segment"] == "Encalm Eats"]
    aug = ee[ee["month"] == 8]
    # heads 3 + 4 + 1 = 8 lakhs; children sum to the same, less CDH (2 L, has a note)
    assert aug["aop"].sum() == pytest.approx(6 * 100000)
    outlets = set(aug["outlet"])
    assert "Chaayos (T3 Arrivals)" in outlets              # not merged into Gate 54
    assert "Chaayos T3 Gate No 54" in outlets
    assert "Starbucks (T3 Departures) #2" in outlets       # duplicate kept separate
    assert "Event/ODC" in outlets
    assert set(aug[aug["outlet"] == "Event/ODC"]["location"]) == {"Delhi"}
    assert not ee.duplicated(["location", "outlet", "year", "month"]).any()
    assert res["warnings"] == []


def test_lines_with_a_note_are_not_imported():
    res = esp.parse_eats_and_skyplates_aop(_workbook())
    assert not res["aop_rows"]["outlet"].str.startswith("Cafe Delhi Heights").any()
    m = res["mapping"]
    cdh = m[m["aop_name"] == "CDH"].iloc[0]
    assert cdh["result"].startswith("Not imported") and cdh["note"] == "some note"
    assert (m[m["aop_name"] != "CDH"]["result"] == esp.IMPORTED).all()


def _bhogapuram_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 4, 100000)
    for i, h in enumerate(["Geographical Segment", "Business Segment", "Unit-ID"]):
        ws.cell(2, 2 + i, h)
    ws.cell(2, 6, dt.datetime(2026, 9, 1))
    rows = [("Bhogapuram", "F&B", "Airport Outlets F&B", 10),
            (None, None, "IndiSip", 1),
            (None, "Change as per actual", "Starbucks", 2),
            (None, None, "Starbucks", 3),
            (None, None, "IndiSip", 4)]
    for r, (geo, seg, unit, v) in enumerate(rows, start=3):
        ws.cell(r, 2, geo); ws.cell(r, 3, seg); ws.cell(r, 4, unit); ws.cell(r, 6, v)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def test_repeated_names_map_by_occurrence_and_combined_line():
    rows = esp.parse_eats_and_skyplates_aop(_bhogapuram_workbook())["aop_rows"]
    by_outlet = dict(zip(rows["outlet"], rows["aop"] / 100000))
    assert by_outlet == {"IndiSip Café": 1, "Indisip Bistro": 4,
                         "Starbucks Departure + Starbucks Arrival": 3}


def test_combined_aop_split_by_revenue_share():
    import pandas as pd
    aop = pd.DataFrame({"location": ["Bhogapuram"] * 2, "business_unit": ["Airport Outlets"] * 2,
                        "outlet": ["Starbucks Departure + Starbucks Arrival", "KFC (Airport Outlets)"],
                        "aop": [300.0, 50.0]})
    act = pd.DataFrame({"outlet": ["Starbucks Departure", "Starbucks Arrival", "Svara"],
                        "revenue": [200.0, 100.0, 10.0]})
    out = esp.align_aop_to_outlets(aop, act).set_index("outlet")["aop"]
    assert out["Starbucks Departure"] == pytest.approx(200)
    assert out["Starbucks Arrival"] == pytest.approx(100)
    assert out["KFC (Airport Outlets)"] == 50 and out.sum() == pytest.approx(350)
    # no sales at either store yet: the line stays whole (shown as a planned outlet)
    none = esp.align_aop_to_outlets(aop, pd.DataFrame({"outlet": ["Svara", "KFC (Airport Outlets)"],
                                                       "revenue": [1.0, 2.0]}))
    assert "Starbucks Departure + Starbucks Arrival" in set(none["outlet"])


def test_mismatch_is_reported():
    buf = _workbook()
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 4).value == "T3 Arrivals F&B":
            ws.cell(r, 8, 99)                              # August head no longer = children
    out = io.BytesIO(); wb.save(out); out.seek(0)
    res = esp.parse_eats_and_skyplates_aop(out)
    assert any("T3 Arrivals" in w for w in res["warnings"])


def test_align_single_outlet():
    import pandas as pd
    aop = pd.DataFrame({"location": ["Delhi", "Delhi"], "business_unit": ["x", "y"],
                        "outlet": ["A (T3)", "B (T3)"], "aop": [1.0, 2.0]})
    act = pd.DataFrame({"outlet": ["Encalm Eats"]})
    out = esp.align_aop_to_outlets(aop, act)
    assert list(out["outlet"]) == ["Encalm Eats"] and out["aop"].sum() == 3.0
    act2 = pd.DataFrame({"outlet": ["A (T3)", "Z"]})
    assert esp.align_aop_to_outlets(aop, act2) is aop


def test_la22_stored_once_after_canonicalization():
    from modules import aop_parser
    from modules.database import canonicalize_outlet_name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 4, 100000)
    for i, h in enumerate(["Geographical Segment", "Business Segment", "Unit-ID", "Sales Unit-ID"]):
        ws.cell(2, 2 + i, h)
    ws.cell(2, 6, dt.datetime(2026, 8, 3))
    ws.cell(3, 2, "Delhi"); ws.cell(3, 3, "Lounge"); ws.cell(3, 4, "LA22"); ws.cell(3, 6, 33.9)
    ws.cell(4, 2, "Bhogapuram"); ws.cell(4, 3, "Lounge"); ws.cell(4, 4, "Domestic Card Lounge"); ws.cell(4, 6, 60)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    df = aop_parser.parse_outlet_monthly_aop(buf)["aop_rows"]
    la = df[df["location"] == "Delhi"].copy()
    la["outlet"] = la["outlet"].map(canonicalize_outlet_name)
    assert la["aop"].sum() == pytest.approx(33.9 * 100000)      # was 3x before the fix
    assert df[df["location"] == "Bhogapuram"]["aop"].sum() == pytest.approx(60 * 100000)
