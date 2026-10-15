"""Traffic workbook in the Sep-2026 layout: Delhi day grid with the month as a
date in A1, Hyderabad / Vizag side-by-side Domestic | International, and a
warning for sheets that give no rows."""

import datetime as dt
import io

import openpyxl
import pytest

from modules import terminal_mapping as tm
from modules import traffic_parser as tp

DAYS = 3


def _workbook(extra_sheet: str | None = None) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delhi"
    ws.append([dt.datetime(2026, 9, 1)])                       # month as a real date
    ws.append(["Day", "T3 International", None, "T3 Domestic", None,
               "T2 Domestic", None, "Terminal 1", None, "G.Total"])
    ws.append([None, "Dep.", "Arr.", "Dep.", "Arr.", "Dep.", "Arr.", "Dep.", "Arr.", None])
    for d in range(1, DAYS + 1):
        ws.append([d, 100, 90, 80, 70, 60, 50, 40, 30, "  520"])
    ws.append(["Total", "300", "270", "240", "210", "180", "150", "120", "90", "1,560"])

    for name in ("Hyd", "Vizag"):
        s = wb.create_sheet(name)
        s.append(["DOMESTIC PAX (Schedule)", None, None, "International PAX (Schedule)", None, None])
        s.append([None, "ARRIVAL", "DEPARTURE", "TOTAL", "ARRIVAL", "DEPARTURE", "TOTAL"])
        for d in range(1, DAYS + 1):
            s.append([dt.datetime(2026, 9, d), 10, 20, 30, 1, 2, 3])
        s.append(["Total", 30, 60, 90, 3, 6, 9])

    if extra_sheet:
        s = wb.create_sheet(extra_sheet)
        s.append(["Something", "else"])
        s.append(["not", "traffic"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _totals(df):
    return df.groupby(["location", "terminal"])["traffic"].sum().to_dict()


def test_delhi_day_grid_uses_date_in_a1_as_month():
    df = tp.parse_traffic_auto(_workbook(), "sept.xlsx")
    delhi = df[df["location"] == "Delhi"]
    assert len(delhi) == 8 * DAYS
    assert {d.month for d in delhi["date"]} == {9} and {d.year for d in delhi["date"]} == {2026}
    t = _totals(delhi)
    assert t[("Delhi", "T3 Int Dep")] == 300 and t[("Delhi", "T1 Arr")] == 90


def test_vizag_sheet_is_bhogapuram_and_hyd_stays_hyderabad():
    df = tp.parse_traffic_auto(_workbook(), "sept.xlsx")
    t = _totals(df)
    for loc in ("Hyderabad", "Bhogapuram"):
        assert t[(loc, "Domestic")] == 30 * DAYS        # arrival + departure, not TOTAL
        assert t[(loc, "International")] == 3 * DAYS
    assert len(df) == 8 * DAYS + 2 * 2 * DAYS
    assert df.attrs["skipped_sheets"] == []


def test_unrecognised_sheet_is_reported():
    df = tp.parse_traffic_auto(_workbook(extra_sheet="Notes"), "sept.xlsx")
    assert any("'Notes'" in m for m in df.attrs["skipped_sheets"])


def test_airport_sheet_with_no_rows_is_reported():
    buf = _workbook()
    wb = openpyxl.load_workbook(buf)
    wb.create_sheet("Goa").append(["just a title"])
    out = io.BytesIO(); wb.save(out); out.seek(0)
    df = tp.parse_traffic_auto(out, "sept.xlsx")
    assert any("'Goa' (Goa)" in m for m in df.attrs["skipped_sheets"])


@pytest.mark.parametrize("outlet,terminal", [
    ("Domestic Lounge T3", "Domestic"),
    ("International Lounge", "International"),
    ("Reserved Lounge", "International"),
    ("Atithya", "All"),
    ("Porter", "All"),
    ("Baggage Wrapping", "All"),
    ("Domestic Spa - T3", tm.MAIN_TERMINAL),     # not listed → whole airport
])
def test_bhogapuram_terminal_mapping(outlet, terminal):
    assert tm.get_terminal_for_outlet(outlet, "Bhogapuram") == terminal


def test_may26_hyd_layout_heading_above_departure_column():
    """'Domestic' / 'International' sit right above the Departure columns
    (Departure before Arrival). The International heading must not make
    its Departure column count as Domestic."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HYD Traffic May-26"
    ws.append([None, "Domestic", None, "International", None])
    ws.append(["Date", "Departure", "Arrival", "Departure", "Arrival"])
    for d in range(1, DAYS + 1):
        ws.append([dt.datetime(2026, 5, d), 100, 200, 10, 20])
    ws.append(["Total", 300, 600, 30, 60])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    t = _totals(tp.parse_traffic_auto(buf, "may.xlsx"))
    assert t[("Hyderabad", "Domestic")] == 300 * DAYS
    assert t[("Hyderabad", "International")] == 30 * DAYS


@pytest.mark.parametrize("measures", [("PAX", "Flights"), ("ATMs", "PAX")])
def test_hyd_layout_reads_pax_not_flights(measures):
    """Apr-26 template: per side PAX Arr/Dep then Flights Arr/Dep; older
    template: ATMs before PAX. Only the PAX columns are traffic."""
    first, second = measures
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hyd Traf"
    ws.append([dt.datetime(2026, 4, 1)])
    ws.append([None, "Domestic", None, None, None, "International", None, None, None])
    ws.append([None, first, None, second, None, first, None, second, None])
    ws.append(["Date"] + ["Arrival", "Departure"] * 4)
    pax_dom, pax_int, small = (100, 200), (10, 20), (2, 3)
    for d in range(1, DAYS + 1):
        dom = (pax_dom + small) if first == "PAX" else (small + pax_dom)
        intl = (pax_int + small) if first == "PAX" else (small + pax_int)
        ws.append([dt.datetime(2026, 4, d), *dom, *intl])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    t = _totals(tp.parse_traffic_auto(buf, "apr.xlsx"))
    assert t[("Hyderabad", "Domestic")] == 300 * DAYS
    assert t[("Hyderabad", "International")] == 30 * DAYS


def test_goa_two_month_blocks_keeps_only_current_month():
    """Jul-26 Goa sheet: left block = July 2025 (daily figures), right block
    = July 2026. Only July 2026 is traffic for this file, and the two blocks
    must never be added together."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gox Traffic July 25 July26"
    block_head = [["Domestic", None, None, None, None, "International", None, None, None],
                  ["Date", "PAX", None, "Flights", None, "PAX", None, "Flights", None],
                  [None, "Arrival", "Departure", "Arrival", "Departure", "Arrival", "Departure", "Arrival", "Departure"]]
    ws.append([dt.datetime(2025, 7, 25)] + [None] * 9 + [dt.datetime(2026, 7, 1)])
    for h in block_head:
        ws.append(h + [None] + h)
    for d in range(1, DAYS + 1):
        left = [dt.datetime(2025, 7, d), 1000, 1000, 9, 9, 50, 50, 1, 1]
        right = [dt.datetime(2026, 7, d), 300, 400, 3, 3, 20, 30, 1, 1]
        ws.append(left + [None] + right)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    df = tp.parse_traffic_auto(buf, "jul.xlsx")
    goa = df[df["location"] == "Goa"]
    assert {(d.year, d.month) for d in goa["date"]} == {(2026, 7)}
    t = _totals(goa)
    assert t[("Goa", "Domestic")] == 700 * DAYS
    assert t[("Goa", "International")] == 50 * DAYS


def test_monthly_cross_tab_stores_domestic_international_split():
    """Monthly 'Apr24-Mar26' sheets: Hyderabad / Goa are stored as Domestic
    and International (dep + arr), not one airport total, so lounge PEN %
    uses real figures. Delhi keeps its terminal pools."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HYD"
    ws.append(["HYD"])
    ws.append(["Month", "International", None, "Domestic", None, "Sub Totals", None, "G.Total "])
    ws.append([None, "Dep.", "Arr.", "Dep.", "Arr.", "Dep.", "Arr.", None])
    ws.append(["Apr'24", 10, 20, 100, 200, 110, 220, 330])
    ws.append(["May'24", 11, 21, "1,000", "2,000", 1011, 2021, 3032])     # Indian-format text
    ws.append(["Fy24-25", 21, 41, 1100, 2200, 1121, 2241, 3362])          # rollup, skipped
    g = wb.create_sheet("GOX")
    g.append(["GOX"])
    g.append(["Month", "Domestic", None, "International", None, "Sub Totals", None, "G.Total "])
    g.append([None, "Arr.", "Dep.", "Arr.", "Dep.", "Dep.", "Arr.", None])
    g.append(["Apr'24", 300, 400, 5, 6, 406, 305, 711])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    df = tp.parse_traffic_auto(buf, "monthly.xlsx")
    assert set(df["granularity"]) == {"monthly"}
    t = df.assign(m=[d.month for d in df["date"]]).groupby(["location", "m", "terminal"])["traffic"].sum().to_dict()
    assert t[("Hyderabad", 4, "International")] == 30 and t[("Hyderabad", 4, "Domestic")] == 300
    assert t[("Hyderabad", 5, "Domestic")] == 3000
    assert t[("Goa", 4, "Domestic")] == 700 and t[("Goa", 4, "International")] == 11
    assert "" not in set(df["terminal"])
