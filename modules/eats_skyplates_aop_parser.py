"""
eats_skyplates_aop_parser.py — AOP importer for Sky Plates (IFK) and
Encalm Eats (F&B), read from the same AOP master workbook as aop_parser.py.

aop_parser.py imports the EHPL rows (Lounge / SPA / Atithya / Others, etc.)
and deliberately skips IFK and F&B. This module picks those up.

HOW THE SHEET IS READ (label-driven — no hard-coded row numbers)
────────────────────────────────────────────────────────────────
The header row is found by its "Geographical Segment" / "Business Segment" /
"Unit-ID" labels. Every datetime cell in that row is a monthly column
(both the FY25-26 and FY26-27 blocks are read). The units note (100000 =
lakhs) is read from the rows above the header.

Rows below the header are walked top to bottom:

  • Section head — Geographical + Business Segment both filled, and the
    segment is IFK or F&B. Anything else with both filled closes the section.
  • Child row  — Geographical Segment blank, Unit-ID filled, directly under
    an F&B head. (Column C on child rows sometimes holds a note such as
    "Put in as other rev"; it is kept as a note, never as a segment.)
  • A blank Unit-ID row closes the section.

IFK heads are single rows → one Sky Plates target per location, stored under
outlet "Encalm Sky Plates" (the Sub-Business name used by revenue uploads).

F&B heads are subtotals of their children (verified: equal every month for
every section). Only the CHILDREN are stored, so any consumer that sums
aop_target rows for a location gets the right total with no double count.
An F&B head with no children (Hyderabad "Multiple Outlets") is stored as a
single "Encalm Eats" row. Any head-vs-children mismatch is reported in
`warnings` rather than silently trusted.

OUTLET NAMES
────────────
AOP Unit-IDs are mapped to the outlet names the Encalm Eats DSR uses, keyed
by (section, AOP name) so the same short name in two terminals ("Chaayos" in
T3 Departures vs T3 Arrivals) can't collide. Only mappings that are clear
from the DSR's own outlet list are made. Outlets with no DSR counterpart keep
their AOP name with the section in brackets, e.g. "BAR SHA (T3 Departures)",
so they are never merged into a different outlet's target. They still count
in the location total. `mapping` in the result lists every decision.

Location "All" (Events & ODCs) is stored under Delhi, because the DSR books
Event/ODC revenue under Delhi. Change _LOCATION_MAP to move it.
"""

from __future__ import annotations

import datetime as dt
from typing import Optional

import pandas as pd

from . import database

OUTPUT_COLS = ["segment", "business_unit", "location", "outlet", "year", "month", "aop"]

_HEADER_LABELS = ("geographical segment", "business segment", "unit-id")

_LOCATION_MAP = {
    "delhi": "Delhi", "del": "Delhi",
    "hyd": "Hyderabad", "hyderabad": "Hyderabad",
    "bhogapuram": "Bhogapuram", "bpm": "Bhogapuram", "vtz": "Bhogapuram",
    "goa": "Goa",
    "all": "Delhi",   # Events & ODCs — DSR books Event/ODC under Delhi
}

_SECTION_SEGMENTS = {"ifk": "Sky Plates", "f&b": "Encalm Eats"}

SKY_PLATES_OUTLET = "Encalm Sky Plates"
EATS_AGGREGATE_OUTLET = "Encalm Eats"

# (section label, AOP Unit-ID) -> DSR outlet name.  Section label is the
# head's Unit-ID with a trailing " F&B" removed (see _section_label).
_DSR_OUTLET_MAP: dict[tuple[str, str], str] = {
    ("T3 Departures", "chaayos 42a"):      "Chaayos T3 Gate No. 42A",
    ("T3 Departures", "chaayos 54-55"):    "Chaayos T3 Gate No 54",
    ("T3 Departures", "indisip 44"):       "IndiSip T3 - Mezzanine Gate 44",
    ("T3 Departures", "indisip 3"):        "IndiSip T3 - Int. IFK 02, Gate 03",
    ("T2 Departures", "indisip sha"):      "IndiSip T2 - Bus Gate",
    ("T2 Departures", "indisip prechk"):   "IndiSip T2 - Pre Checkin",
    ("T2 Departures", "haldiram prechk"):  "Haldiram T2 Pre Check-in",
    ("Events & ODCs", "events & odcs"):    "Event/ODC",
}


class EatsAOPParseError(Exception):
    """Raised when the AOP workbook can't be parsed for Eats / Sky Plates."""


# ──────────────────────────────────────────────────────────────────────────────
# Sheet discovery helpers
# ──────────────────────────────────────────────────────────────────────────────

def _norm(v) -> str:
    return " ".join(str(v).split()).strip().lower() if v is not None else ""


def _find_layout(ws) -> Optional[dict]:
    """Locate header row + label columns + monthly date columns in a sheet."""
    for r in range(1, min(ws.max_row, 15) + 1):
        labels = {_norm(ws.cell(r, c).value): c for c in range(1, min(ws.max_column, 20) + 1)}
        if all(l in labels for l in _HEADER_LABELS):
            month_cols = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, dt.datetime):
                    month_cols.append((c, v.year, v.month))
            return {
                "header_row": r,
                "geo_col": labels["geographical segment"],
                "seg_col": labels["business segment"],
                "unit_col": labels["unit-id"],
                "month_cols": month_cols,
            }
    return None


def _read_units(ws, header_row: int) -> float:
    """Units note (e.g. 100000 = lakhs) sits in the rows above the header."""
    for r in range(1, header_row):
        for c in range(1, min(ws.max_column, 10) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v in (100000, 10000000):
                return float(v)
    return 1.0


def _num(v) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _section_label(unit_id: str) -> str:
    label = " ".join(str(unit_id).split())
    if label.lower().endswith(" f&b"):
        label = label[:-4].strip()
    if label.lower() == "multiple outlets":
        label = "Airport Outlets"
    return label


# ──────────────────────────────────────────────────────────────────────────────
# Parser
# ──────────────────────────────────────────────────────────────────────────────

def parse_eats_and_skyplates_aop(file_obj) -> dict:
    """
    Returns:
        aop_rows          DataFrame[segment, business_unit, location, outlet, year, month, aop]
                          (aop in INR, already multiplied by `units`)
        mapping           DataFrame — one line per imported AOP row: section, AOP name,
                          dashboard outlet, matched_to_dsr, FY26-27 total (lakhs)
        warnings          list[str] — head/children mismatches, unknown locations
        units             float
        sky_plates_count  int (rows in aop_rows)
        encalm_eats_count int (rows in aop_rows)
    """
    import openpyxl

    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception as exc:
        raise EatsAOPParseError(f"Cannot open workbook: {exc}") from exc

    ws, layout = None, None
    for name in wb.sheetnames:
        lay = _find_layout(wb[name])
        if lay and lay["month_cols"]:
            ws, layout = wb[name], lay
            break
    if ws is None:
        raise EatsAOPParseError(
            "No sheet has the AOP master layout (a header row with 'Geographical "
            "Segment', 'Business Segment', 'Unit-ID' and monthly date columns)."
        )

    units = _read_units(ws, layout["header_row"])
    gcol, scol, ucol = layout["geo_col"], layout["seg_col"], layout["unit_col"]
    month_cols = layout["month_cols"]

    # Formula cells with no cached value come back as None for every month —
    # the symptom of a workbook saved without recalculation / broken links.
    records: list[dict] = []
    mapping: list[dict] = []
    warnings: list[str] = []

    def monthly(row: int) -> dict:
        return {(y, m): _num(ws.cell(row, c).value) for c, y, m in month_cols}

    def fy2627(vals: dict) -> float:
        return sum(v or 0 for (y, m), v in vals.items()
                   if (y == 2026 and m >= 4) or (y == 2027 and m <= 3))

    def emit(segment, bu, location, outlet, vals):
        for (y, m), v in vals.items():
            if v is None or v <= 0:
                continue
            records.append({"segment": segment, "business_unit": bu, "location": location,
                            "outlet": outlet, "year": y, "month": m, "aop": v * units})

    section = None  # dict(kind, location, label, head_row, head_vals, children=[...])

    def close_section():
        nonlocal section
        if section is None:
            return
        s = section
        section = None
        if s["kind"] == "Sky Plates":
            emit("Sky Plates", "IFK", s["location"], SKY_PLATES_OUTLET, s["head_vals"])
            mapping.append({"segment": "Sky Plates", "location": s["location"], "section": "IFK",
                            "aop_name": s["head_name"], "dashboard_outlet": SKY_PLATES_OUTLET,
                            "matched_to_dsr": None, "fy26_27_lakhs": round(fy2627(s["head_vals"]), 1),
                            "excel_row": s["head_row"]})
            return

        children = s["children"]
        if not children:
            emit("Encalm Eats", s["label"], s["location"], EATS_AGGREGATE_OUTLET, s["head_vals"])
            mapping.append({"segment": "Encalm Eats", "location": s["location"], "section": s["label"],
                            "aop_name": s["head_name"], "dashboard_outlet": EATS_AGGREGATE_OUTLET,
                            "matched_to_dsr": False, "fy26_27_lakhs": round(fy2627(s["head_vals"]), 1),
                            "excel_row": s["head_row"]})
            return

        # Reconcile head vs children month by month
        for key, hv in s["head_vals"].items():
            cs = sum((c["vals"].get(key) or 0) for c in children)
            if abs((hv or 0) - cs) > 0.05:
                warnings.append(
                    f"{s['location']} {s['label']} {key[0]}-{key[1]:02d}: section total "
                    f"{(hv or 0):,.2f} L ≠ sum of outlets {cs:,.2f} L (outlet rows used)."
                )

        seen: dict[str, int] = {}
        for c in children:
            key = (s["label"], _norm(c["name"]))
            dsr = _DSR_OUTLET_MAP.get(key)
            outlet = dsr or f"{c['name']} ({s['label']})"
            seen[outlet] = seen.get(outlet, 0) + 1
            if seen[outlet] > 1:   # e.g. two "Starbucks" rows at Bhogapuram
                outlet = f"{outlet} #{seen[outlet]}"
            emit("Encalm Eats", s["label"], s["location"], outlet, c["vals"])
            mapping.append({"segment": "Encalm Eats", "location": s["location"], "section": s["label"],
                            "aop_name": c["name"], "dashboard_outlet": outlet,
                            "matched_to_dsr": bool(dsr), "note": c["note"],
                            "fy26_27_lakhs": round(fy2627(c["vals"]), 1), "excel_row": c["row"]})

    for r in range(layout["header_row"] + 1, ws.max_row + 1):
        geo = ws.cell(r, gcol).value
        seg = ws.cell(r, scol).value
        unit = ws.cell(r, ucol).value
        unit_s = " ".join(str(unit).split()) if unit is not None else ""

        if geo not in (None, "") and seg not in (None, ""):
            close_section()
            kind = _SECTION_SEGMENTS.get(_norm(seg))
            if kind is None:
                continue
            loc = _LOCATION_MAP.get(_norm(geo))
            if loc is None:
                warnings.append(f"Row {r}: unknown location '{geo}' for {seg} '{unit_s}' — skipped.")
                continue
            section = {"kind": kind, "location": loc, "label": _section_label(unit_s),
                       "head_name": unit_s, "head_row": r, "head_vals": monthly(r), "children": []}
            continue

        if not unit_s:
            close_section()
            continue

        if section is not None and section["kind"] == "Encalm Eats" and geo in (None, ""):
            section["children"].append({"name": unit_s, "row": r, "vals": monthly(r),
                                        "note": " ".join(str(seg).split()) if seg else ""})
    close_section()

    df = pd.DataFrame.from_records(records, columns=OUTPUT_COLS) if records \
        else pd.DataFrame(columns=OUTPUT_COLS)
    map_df = pd.DataFrame(mapping)

    if not map_df.empty and map_df["fy26_27_lakhs"].sum() == 0 and df.empty:
        warnings.append(
            "Every IFK / F&B value is blank or zero. If the workbook was saved without "
            "recalculating its linked files, the formula cells have no stored values."
        )

    return {
        "aop_rows": df,
        "mapping": map_df,
        "warnings": warnings,
        "units": units,
        "sky_plates_count": int((df["segment"] == "Sky Plates").sum()) if not df.empty else 0,
        "encalm_eats_count": int((df["segment"] == "Encalm Eats").sum()) if not df.empty else 0,
    }


def save_eats_skyplates_aop(file_obj, source_file: str = "AOP_Revenue_MASTER.xlsx",
                            parsed: Optional[dict] = None) -> dict:
    """Parse (unless `parsed` is given) and insert into aop_target."""
    result = parsed or parse_eats_and_skyplates_aop(file_obj)
    df = result["aop_rows"]
    if df.empty:
        return {**result, "db_inserted": 0, "db_skipped": 0}
    db_result = database.save_aop_targets(df, source_file=source_file)
    return {**result,
            "db_inserted": db_result.get("inserted", 0),
            "db_skipped": db_result.get("skipped", 0)}


# ──────────────────────────────────────────────────────────────────────────────
# Dashboard helpers
# ──────────────────────────────────────────────────────────────────────────────

def load_segment_aop(segment: str, start_date: dt.date, end_date: dt.date,
                     location: Optional[str] = None) -> pd.DataFrame:
    """
    AOP rows for one segment ("Sky Plates" / "Encalm Eats") over a date range,
    prorated to the number of days selected in each month.
    Returns DataFrame[location, business_unit, outlet, aop] (INR), or empty.

    Proration is per month: a range covering 10 of 31 days of August takes
    10/31 of August's target, and a range spanning two months takes each
    month's share separately.
    """
    import calendar

    try:
        raw = database.load_aop_targets_for_range(start_date, end_date)
    except Exception:
        return pd.DataFrame(columns=["location", "business_unit", "outlet", "aop"])
    if raw is None or raw.empty or "segment" not in raw.columns:
        return pd.DataFrame(columns=["location", "business_unit", "outlet", "aop"])

    df = raw[raw["segment"] == segment].copy()
    df["location"] = df["location"].astype(str).str.strip().str.title()
    if location and location != "All Locations":
        df = df[df["location"] == location.strip().title()]
    if df.empty:
        return pd.DataFrame(columns=["location", "business_unit", "outlet", "aop"])

    def share(row):
        y, m = int(row["year"]), int(row["month"])
        dim = calendar.monthrange(y, m)[1]
        first, last = dt.date(y, m, 1), dt.date(y, m, dim)
        lo, hi = max(first, start_date), min(last, end_date)
        days = (hi - lo).days + 1
        return max(days, 0) / dim

    df["aop"] = pd.to_numeric(df["aop"], errors="coerce") * df.apply(share, axis=1)
    df = df[df["aop"].notna() & (df["aop"] > 0)]
    if "business_unit" not in df.columns:
        df["business_unit"] = None
    return (df.groupby(["location", "business_unit", "outlet"], as_index=False, dropna=False)["aop"]
              .sum())


def align_aop_to_outlets(aop_df: Optional[pd.DataFrame],
                         actual_df: Optional[pd.DataFrame]) -> Optional[pd.DataFrame]:
    """
    Make AOP outlet names line up with the outlets actually present in the
    revenue data for ONE location.

    If the actual data has a single outlet (revenue uploads carry one
    "Encalm Eats" / "Encalm Sky Plates" row per location per day) and none
    of the AOP outlets match it, the location's whole AOP is placed on that
    outlet so the table shows target vs actual. Otherwise AOP is returned
    unchanged (outlet-level DSR data matches outlet-level AOP by name).
    """
    if aop_df is None or aop_df.empty:
        return None
    if actual_df is None or actual_df.empty or "outlet" not in actual_df.columns:
        return aop_df
    actual_outlets = set(actual_df["outlet"].astype(str).str.strip())
    if actual_outlets & set(aop_df["outlet"].astype(str)):
        return aop_df
    if len(actual_outlets) == 1:
        only = next(iter(actual_outlets))
        out = aop_df.copy()
        out["outlet"] = only
        return out.groupby(["location", "outlet"], as_index=False)["aop"].sum()
    return aop_df
